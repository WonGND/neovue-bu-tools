"""BU 트렌드 리포트 — 과거 파일 읽기 + 리포트 생성 (화면과 분리).

그래프는 `CHARTS` 모듈을 통해서만 그린다. 엑셀 기본 차트(B안)로 바꿀 때는
같은 함수 이름을 가진 모듈로 CHARTS만 교체하면 된다.
"""
from __future__ import annotations

import io
import os
import re
import sys
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

_HERE = Path(__file__).resolve().parent
for _p in (_HERE.parent / "common", _HERE / "common"):
    if _p.exists() and str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import charts as CHARTS  # noqa: E402
import history_db as db  # noqa: E402
from capability import SPECS, capability, grade, is_constant  # noqa: E402

VERSION = "v0.3"
FONT = "맑은 고딕"
DEFAULT_SETTINGS = {
    "bu_sigma_increase_pct": 20.0,      # 직전 이벤트 대비 BU σ 증가 경고(%)
    "wu_mean_shift": 1.0,               # 평균 이동 경고(%p)
    "bu_mean_shift": 3.0,
    "ng_rate_increase": 2.0,            # NG율 증가 경고(%p)
    "daily_mean_dev_sigma": 0.5,        # 이벤트 안 일자 평균이 이벤트 평균에서 몇 σ 벗어나면 표시
    "white_typ": {"WINE": [0.293, 0.320], "VDA129BLAC01": [0.293, 0.320], "VMA129BLAC01": [0.293, 0.320]},
    "gcs": 0.005,
}
TEMPLATE_COLS = {"LotID": 3, "Red_X": 6, "Red_Y": 7, "Green_X": 8, "Green_Y": 9, "Blue_X": 10, "Blue_Y": 11,
                 "White_x": 13, "White_y": 14, "White_LV": 15, "Black_LV": 16, "WU": 19, "BU": 20, "EEPROM_59": 27}
_thin = Side(style="thin", color="CBD5E1")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_HEAD = PatternFill("solid", fgColor="1E3A5F")


# ------------------------------------------------------------------ 과거 Daily data 파일 읽기
def date_from_filename(path) -> date | None:
    m = re.search(r"(20\d{6})[_-]?\d{0,6}", Path(path).stem)
    if m:
        try:
            return datetime.strptime(m.group(1), "%Y%m%d").date()
        except ValueError:
            return None
    return None


class NotDailyFile(ValueError):
    """Tool A 'Daily data 통합정리' 양식이 아닌 파일"""


def is_daily_sheet(ws) -> bool:
    """양식 표시: B19 'CPK' 머리글 + P2·P3 모델 칸 ('• Cut. Model :', '• Model :')"""
    def txt(cell):
        v = ws[cell].value
        return "" if v is None else str(v).strip()
    return txt("B19").upper() == "CPK" and "MODEL" in txt("P2").upper() and "MODEL" in txt("P3").upper()


def parse_daily_file(path) -> dict:
    """Tool A 'Daily data 통합정리' 결과 파일 → {'rows': DataFrame, 'model': .., 'code': .., 'date': ..}"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    if not is_daily_sheet(ws):                 # CA410/LMK6 cleaned 등 다른 엑셀은 읽지 않음
        wb.close()
        raise NotDailyFile(Path(path).name)
    cut = ws["R2"].value or ""
    code = ws["R3"].value or ""
    recs = []
    for row in ws.iter_rows(min_row=20, max_col=27, values_only=True):
        lot = row[2]
        if lot is None or str(lot).strip() == "":
            if recs:
                break
            continue
        recs.append({k: row[c - 1] for k, c in TEMPLATE_COLS.items()})
    wb.close()
    df = pd.DataFrame(recs, columns=list(TEMPLATE_COLS))
    for c in df.columns:
        if c not in ("LotID", "EEPROM_59"):
            df[c] = pd.to_numeric(df[c], errors="coerce")
    df["CR"] = df["White_LV"] / df["Black_LV"]
    d = date_from_filename(path) or datetime.fromtimestamp(os.path.getmtime(path)).date()
    return {"rows": df, "model": str(cut).strip(), "code": str(code).strip(), "date": d,
            "ld": bool(len(df) > 1 and is_constant(df["BU"]))}


def rows_for_db(parsed: dict, meas_date: date, path, manual_event: str = "") -> pd.DataFrame:
    df = parsed["rows"].copy()
    df["모델명"] = parsed["model"]
    df["모델코드"] = df["LotID"].map(db.model_code).replace("", parsed["code"])
    df["측정일시"] = pd.Timestamp(meas_date)
    df["생산일"] = df["LotID"].map(db.prod_date)
    df["판정"] = np.where((df["WU"] < SPECS["WU"][0]) | ((df["BU"] < SPECS["BU"][0]) & (not parsed["ld"])), "NG", "OK")
    df["LD"] = "Y" if parsed["ld"] else ""
    df["이벤트(수동)"] = manual_event or ""
    df["출처"] = db.UPLOAD_SOURCE
    df["원본파일"] = Path(path).name
    df["Tool"] = f"NEOVUE BU_TREND {VERSION}"
    df["EEPROM_59"] = df["EEPROM_59"].astype(str)
    return df


# ------------------------------------------------------------------ 엑셀 도우미
def _title(ws, cell, text, sub=None):
    ws[cell] = text
    ws[cell].font = Font(name=FONT, size=16, bold=True, color="1E3A5F")
    if sub:
        c = ws[cell]
        ws.cell(c.row + 1, c.column, sub).font = Font(name=FONT, size=9, color="6B7280")


def _table(ws, r0, c0, header, rows, fills=None, widths=None):
    for j, h in enumerate(header):
        c = ws.cell(r0, c0 + j, h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = _HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BOX
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = ws.cell(r0 + i, c0 + j, v)
            c.font = Font(name=FONT, size=9, bold=(j == 0))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = _BOX
            if fills and (i - 1, j) in fills:
                c.fill = PatternFill("solid", fgColor=fills[(i - 1, j)])
    return r0 + len(rows) + 1


def _img(ws, png: bytes, anchor: str, width_px: int) -> int:
    img = XLImage(io.BytesIO(png))
    ratio = width_px / img.width
    img.width, img.height = width_px, int(img.height * ratio)
    ws.add_image(img, anchor)
    return int(img.height * 0.75 / 15) + 1


def _r(v, nd=2):
    return round(float(v), nd) if v is not None and np.isfinite(v) else "N/A"


# ------------------------------------------------------------------ 리포트
def _event_stats(name, g, ev_row):
    g = g.sort_values("측정일시")
    ld = (g["LD"].astype(str) == "Y").any() or (len(g) > 1 and is_constant(g["BU"]))
    cw = capability(g["WU"], "WU")
    cb = capability(g["BU"], "BU", ld=ld)
    ng = ((g["WU"] < SPECS["WU"][0]) | ((g["BU"] < SPECS["BU"][0]) & (not ld))).mean() * 100
    d = pd.to_datetime(g["측정일"], errors="coerce").dropna()
    period = ""
    if len(d):
        a, b = d.min().date(), d.max().date()
        period = str(a) if a == b else f"{a} ~ {b}"
    return {"name": name, "g": g, "ld": ld, "cw": cw, "cb": cb, "ng": ng, "period": period,
            "model": ", ".join(sorted(set(g["모델명"].dropna().astype(str)) - {""})) or (ev_row["모델"] if ev_row is not None else ""),
            "memo": ev_row["메모"] if ev_row is not None else "",
            "lv": g["White_LV"].mean(), "lv_sd": g["White_LV"].std(), "cr": g["CR"].mean(), "cr_sd": g["CR"].std(),
            "wx": g["White_x"].mean(), "wy": g["White_y"].mean()}


# ------------------------------------------------------------------ 변화 감지 (v0.3: 설명 보강, 별도 시트)
CHANGE_MEANING = {
    "BU σ": "BU 산포(흩어짐)가 커짐 → 자재·공정 편차, 측정 조건(지그·암실·설비 상태) 확인",
    "BU 평균": "BU 분포 전체가 이동 → 자재 Lot · 공정 조건 변경 여부 확인",
    "BU Cpk 등급": "규격 하한(50) 대비 여유가 줄어 등급 하락 → σ·평균 중 원인 확인",
    "WU 평균": "WU 분포 전체가 이동 → 자재 Lot · 공정 조건 변경 여부 확인",
    "WU Cpk 등급": "규격 하한(80) 대비 여유가 줄어 등급 하락 → 평균·산포 중 원인 확인",
    "NG율": "규격 미달(WU<80 또는 BU<50) 비율 증가 → 'NG 이력' 시트에서 Lot ID 확인",
    "전체": "모델이 달라 비교하지 않음 (같은 모델끼리만 비교)",
}


def _criteria(st) -> dict:
    return {"BU σ": f"+{st['bu_sigma_increase_pct']:.0f}% 이상 증가",
            "BU 평균": f"±{st['bu_mean_shift']:g} 이상 이동",
            "BU Cpk 등급": "한 단계 이상 하락",
            "WU 평균": f"±{st['wu_mean_shift']:g} 이상 이동",
            "WU Cpk 등급": "한 단계 이상 하락",
            "NG율": f"+{st['ng_rate_increase']:g}%p 이상 증가",
            "전체": "—"}


def _change_detection(stats, st):
    """직전 이벤트(시간순 바로 앞)와 비교. 반환: (행 목록, 확인 필요 건수)
    행: [이벤트, 비교 대상, 항목, 변화, 변화량, 경고 기준, 결과, 의미]"""
    crit = _criteria(st)
    flags, n_hit = [], 0
    order = ["부적합", "주의", "적합", "양호"]
    for prev, cur in zip(stats, stats[1:]):
        if str(prev["model"]).strip().upper() != str(cur["model"]).strip().upper():
            flags.append([cur["name"], prev["name"], "전체", f"{prev['model']} → {cur['model']}", "", crit["전체"],
                          "비교 생략", CHANGE_MEANING["전체"]])
            continue
        checks = []
        if not (prev["ld"] or cur["ld"]):
            d_sig = (cur["cb"].sigma_overall / prev["cb"].sigma_overall - 1) * 100
            checks.append(("BU σ", f"{prev['cb'].sigma_overall:.2f} → {cur['cb'].sigma_overall:.2f}", f"{d_sig:+.0f}%",
                           d_sig >= st["bu_sigma_increase_pct"]))
            dm = cur["cb"].mean - prev["cb"].mean
            checks.append(("BU 평균", f"{prev['cb'].mean:.2f} → {cur['cb'].mean:.2f}", f"{dm:+.2f}",
                           abs(dm) >= st["bu_mean_shift"]))
            g0, g1 = grade(prev["cb"].cpk)[0], grade(cur["cb"].cpk)[0]
            checks.append(("BU Cpk 등급", f"{g0} → {g1}", f"{cur['cb'].cpk - prev['cb'].cpk:+.2f}",
                           order.index(g1) < order.index(g0)))
        dm = cur["cw"].mean - prev["cw"].mean
        checks.append(("WU 평균", f"{prev['cw'].mean:.2f} → {cur['cw'].mean:.2f}", f"{dm:+.2f}",
                       abs(dm) >= st["wu_mean_shift"]))
        g0, g1 = grade(prev["cw"].cpk)[0], grade(cur["cw"].cpk)[0]
        checks.append(("WU Cpk 등급", f"{g0} → {g1}", f"{cur['cw'].cpk - prev['cw'].cpk:+.2f}",
                       order.index(g1) < order.index(g0)))
        dng = cur["ng"] - prev["ng"]
        checks.append(("NG율", f"{prev['ng']:.1f}% → {cur['ng']:.1f}%", f"{dng:+.1f}%p", dng >= st["ng_rate_increase"]))
        for item, chg, amt, hit in checks:
            n_hit += bool(hit)
            flags.append([cur["name"], prev["name"], item, chg, amt, crit[item], "⚠ 확인 필요" if hit else "정상",
                          CHANGE_MEANING[item] if hit else ""])
    return flags, n_hit


def _change_sheet(wb, flags, st):
    ws = wb.create_sheet("변화 감지", 1)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJ", (2, 13, 13, 13, 19, 10, 17, 13, 64, 2)):
        ws.column_dimensions[col].width = w
    _title(ws, "B2", "변화 감지 — 직전 이벤트 대비", "이벤트 요약 시트의 이벤트를 시간 순으로 놓고, 각 이벤트를 바로 앞 이벤트와 비교")
    guide = [("읽는 법", True),
             ("① 비교 방법  각 이벤트를 시간 순으로 바로 앞 이벤트와 비교합니다 (같은 모델끼리만, 모델이 다르면 비교 생략).", False),
             ("② 결과  '⚠ 확인 필요' = 아래 경고 기준을 넘은 변화 → 원인 확인 권장 · '정상' = 기준 이내.  "
              "※ 불량 판정이 아니라 '달라졌다'는 신호입니다.", False),
             ("③ 의미 · 확인할 것  ⚠ 행에만 표시 — 무엇이 변했고 어디를 먼저 볼지.", False),
             ("④ Cpk 등급  양호 ≥ 1.67 · 적합 ≥ 1.33 · 주의 ≥ 1.00 · 부적합 < 1.00 (Cp·Cpk 같은 기준).  "
              "규격: WU 80~100 · BU 50~100", False)]
    fill = PatternFill("solid", fgColor="EFF6FF")
    for i, (txt, bold) in enumerate(guide):
        r = 4 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        c = ws.cell(r, 2, txt)
        c.font = Font(name=FONT, size=10 if not bold else 11, bold=bold, color="1E3A5F" if bold else "374151")
        c.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        for cc in range(2, 10):
            ws.cell(r, cc).fill = fill
        ws.row_dimensions[r].height = 20
    r = 10
    if flags:
        ff = {}
        for i, f in enumerate(flags):
            if f[6].startswith("⚠"):
                ff[(i, 6)] = "FEE2E2"
                ff[(i, 7)] = "FEF2F2"
            elif f[6] == "비교 생략":
                ff[(i, 6)] = "E5E7EB"
        end = _table(ws, r, 2, ["이벤트", "비교 대상\n(직전 이벤트)", "항목", "변화 (직전 → 이번)", "변화량", "경고 기준",
                               "결과", "의미 · 확인할 것"], flags, ff)
        for rr in range(r + 1, end):
            ws.cell(rr, 9).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.auto_filter.ref = f"B{r}:I{end - 1}"
        ws.freeze_panes = f"B{r + 1}"
        r = end + 1
    else:
        ws.cell(r, 2, "비교할 이벤트가 2개 이상일 때 표시됩니다.").font = Font(name=FONT, size=9, color="6B7280")
        r += 2
    ws.cell(r, 2, "경고 기준 바꾸기").font = Font(name=FONT, size=11, bold=True, color="1E3A5F")
    rows = [["BU σ 증가", f"{st['bu_sigma_increase_pct']:g} %", "bu_sigma_increase_pct"],
            ["BU 평균 이동", f"{st['bu_mean_shift']:g}", "bu_mean_shift"],
            ["WU 평균 이동", f"{st['wu_mean_shift']:g}", "wu_mean_shift"],
            ["NG율 증가", f"{st['ng_rate_increase']:g} %p", "ng_rate_increase"]]
    _table(ws, r + 1, 2, ["항목", "현재 기준", "설정 키 (BU_TREND_settings.json → report)"], rows)
    ws.merge_cells(start_row=r + 1, start_column=4, end_row=r + 1, end_column=5)
    for rr in range(r + 2, r + 2 + len(rows)):
        ws.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=5)


def generate(db_path, selected_events: list[str], out_path, metrics=("BU·WU", "White 휘도·CR", "White x/y"),
             settings: dict | None = None, log=print) -> str:
    st = {**DEFAULT_SETTINGS, **(settings or {})}
    data, events = db.load(db_path)
    if data.empty:
        raise ValueError("누적 DB에 데이터가 없습니다.")
    data, events, _ = db.migrate_period_events(data, events)      # v0.1 기간 이벤트도 같은 방식으로
    data = db.assign_events(data, events)
    first = (data[data["이벤트"] != ""].assign(_d=pd.to_datetime(data["측정일"], errors="coerce"))
             .groupby("이벤트")["_d"].min().sort_values())
    ev_order = [str(e) for e in first.index]
    selected = [e for e in ev_order if e in selected_events] + [e for e in selected_events if e not in ev_order]
    data = data[data["이벤트"].isin(selected)].copy()
    if data.empty:
        raise ValueError("선택한 이벤트에 해당하는 데이터가 없습니다. 날짜를 이벤트로 묶었는지 확인하세요.")
    for c in ("WU", "BU", "White_LV", "Black_LV", "CR", "White_x", "White_y"):
        data[c] = pd.to_numeric(data[c], errors="coerce")
    # 같은 이벤트 안 동일 LotID는 최신 측정 1건
    data = data.sort_values("측정일시").drop_duplicates(subset=["이벤트", "LotID"], keep="last")
    ev_rows = {str(r["이벤트명"]): r for _, r in events.iterrows()}
    stats = [_event_stats(e, data[data["이벤트"] == e], ev_rows.get(e)) for e in selected if (data["이벤트"] == e).any()]
    labels = [s["name"] for s in stats]
    log(f"리포트 대상: {', '.join(labels)} (총 {len(data)}행)")

    wb = Workbook()
    # ---- 1. 이벤트 요약 + 변화 감지
    ws = wb.active
    ws.title = "이벤트 요약"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMNOPQRS", (2, 14, 22, 12, 7, 9, 8, 8, 9, 8, 8, 8, 10, 9, 9, 9, 8, 24, 2)):
        ws.column_dimensions[col].width = w
    _title(ws, "B2", "트렌드 리포트 — 이벤트 요약",
           f"NEOVUE BU_TREND {VERSION} · {datetime.now():%Y-%m-%d %H:%M} · 이벤트 안 동일 LotID는 최신 1건 · Cp/Cpk 같은 기준 색상")
    hdr = ["이벤트", "기간", "모델", "n", "WU 평균", "WU Cp", "WU Cpk", "BU 평균", "BU σ", "BU Cp", "BU Cpk",
           "W 휘도 평균", "CR 평균", "W x", "W y", "NG율", "메모"]
    rows, fills = [], {}
    for i, s in enumerate(stats):
        cw, cb = s["cw"], s["cb"]
        rows.append([s["name"], s["period"], s["model"], len(s["g"]), _r(cw.mean), _r(cw.cp), _r(cw.cpk),
                     "N/A" if s["ld"] else _r(cb.mean), "N/A" if s["ld"] else _r(cb.sigma_overall),
                     "N/A (LD)" if s["ld"] else _r(cb.cp), "N/A (LD)" if s["ld"] else _r(cb.cpk),
                     _r(s["lv"], 1), _r(s["cr"], 0), _r(s["wx"], 4), _r(s["wy"], 4), f"{s['ng']:.1f}%", s["memo"]])
        for j, v in ((5, cw.cp), (6, cw.cpk)) + (() if s["ld"] else ((9, cb.cp), (10, cb.cpk))):
            fills[(i, j)] = grade(v)[1]
    r = _table(ws, 5, 2, hdr, rows, fills)
    ws.auto_filter.ref = f"B5:R{r - 1}"
    r += 1
    flags, n_hit = _change_detection(stats, st)
    ws.cell(r, 2, "변화 감지").font = Font(name=FONT, size=11, bold=True, color="1E3A5F")
    if flags:
        msg = (f"직전 이벤트 대비 확인 필요 {n_hit}건 → '변화 감지' 시트" if n_hit
               else "직전 이벤트 대비 모든 항목이 기준 이내 → '변화 감지' 시트")
    else:
        msg = "비교할 이벤트가 2개 이상일 때 '변화 감지' 시트에 표시됩니다."
    c = ws.cell(r, 4, msg)
    c.font = Font(name=FONT, size=10, bold=bool(n_hit), color="B91C1C" if n_hit else "15803D", underline="single")
    c.hyperlink = "#'변화 감지'!A1"
    _change_sheet(wb, flags, st)

    # ---- 2. 트렌드(거시)
    ws2 = wb.create_sheet("트렌드(거시)")
    ws2.sheet_view.showGridLines = False
    _title(ws2, "B2", "트렌드 그래프 — 이벤트 간 비교 (거시)", "박스플롯: 상자=사분위, 선=1.5×IQR, 점=이상치 · 최신 이벤트 파란색")
    row = 5
    if "BU·WU" in metrics:
        bu_groups = {s["name"]: s["g"]["BU"].dropna() for s in stats if not s["ld"]}
        if bu_groups:
            _img(ws2, CHARTS.boxplot_groups(bu_groups, "BU", SPECS["BU"][0], "BU 분포 (이벤트별)"), f"B{row}", 600)
        _img(ws2, CHARTS.boxplot_groups({s["name"]: s["g"]["WU"].dropna() for s in stats}, "WU", SPECS["WU"][0],
                                        "WU 분포 (이벤트별)"), f"L{row}", 600)
        row += 24
        series = {"WU Cp": [s["cw"].cp for s in stats], "WU Cpk": [s["cw"].cpk for s in stats],
                  "BU Cp": [np.nan if s["ld"] else s["cb"].cp for s in stats],
                  "BU Cpk": [np.nan if s["ld"] else s["cb"].cpk for s in stats]}
        _img(ws2, CHARTS.cp_cpk_bars(labels, series, "Cp · Cpk 추이 (이벤트별)"), f"B{row}", 600)
        _img(ws2, CHARTS.ng_rate_line(labels, [s["ng"] for s in stats], "NG율 추이 (WU<80 또는 BU<50)"), f"L{row}", 600)
        row += 24
    if "White x/y" in metrics:
        models = [str(s["model"]) for s in stats]
        model = max(set(models), key=models.count) if models else ""
        typ = None
        for key, v in st["white_typ"].items():
            if key and key.upper() in model.upper():
                typ = v
        title = "White 색좌표 평균 이동 (GCS 원 안이면 OK)" if typ else f"White 색좌표 평균 이동 (모델 '{model}' Typ 미설정 → 원 생략)"
        _img(ws2, CHARTS.white_xy(labels, [s["wx"] for s in stats], [s["wy"] for s in stats],
                                  tuple(typ) if typ else None, st["gcs"], title), f"B{row}", 480)
    if "White 휘도·CR" in metrics:
        _img(ws2, CHARTS.mean_band(labels, [s["lv"] for s in stats], [s["lv_sd"] for s in stats],
                                   "White 휘도 평균 ± 3σ", "cd/m²"), f"L{row}", 600)
        row += 22
        _img(ws2, CHARTS.mean_band(labels, [s["cr"] for s in stats], [s["cr_sd"] for s in stats],
                                   "CR 평균 ± 3σ", "CR"), f"L{row}", 600)

    # ---- 3. 측정일별(미시)
    ws3 = wb.create_sheet("측정일별(미시)")
    ws3.sheet_view.showGridLines = False
    _title(ws3, "B2", "측정일별 비교 — 이벤트 안 (미시)",
           f"일자 평균이 이벤트 평균에서 {st['daily_mean_dev_sigma']}σ 이상 벗어나면 표시 · 당일 컨디션 편차 확인용")
    row = 5
    for s in stats:
        g = s["g"].copy()
        g["d"] = pd.to_datetime(g["측정일"]).dt.strftime("%m-%d")
        ws3.cell(row, 2, f"■ {s['name']}  ({s['period']})" + (f"  메모: {s['memo']}" if s["memo"] else "")).font = \
            Font(name=FONT, size=11, bold=True, color="1E3A5F")
        row += 1
        days = sorted(g["d"].unique())
        if not s["ld"]:
            _img(ws3, CHARTS.boxplot_groups({d: g.loc[g.d == d, "BU"].dropna() for d in days}, "BU", SPECS["BU"][0],
                                            "BU 측정일별", width=5.0, highlight_last=False), f"B{row}", 430)
        _img(ws3, CHARTS.boxplot_groups({d: g.loc[g.d == d, "WU"].dropna() for d in days}, "WU", SPECS["WU"][0],
                                        "WU 측정일별", width=5.0, highlight_last=False), f"I{row}", 430)
        daily = []
        for d in days:
            gd = g[g.d == d].sort_values("측정일시")
            daily.append((d, len(gd), capability(gd["WU"], "WU"), capability(gd["BU"], "BU", ld=s["ld"])))
        _img(ws3, CHARTS.daily_cpk_line(days, [x[2].cpk for x in daily], [np.nan if s["ld"] else x[3].cpk for x in daily],
                                        s["memo"], "측정일별 Cp · Cpk", wu_cp=[x[2].cp for x in daily],
                                        bu_cp=None if s["ld"] else [x[3].cp for x in daily]), f"P{row}", 460)
        row += 18
        k = st["daily_mean_dev_sigma"]
        trows, tf = [], {}
        for i, (d, n, cw, cb) in enumerate(daily):
            dev_w = abs(cw.mean - s["cw"].mean) >= k * s["cw"].sigma_overall
            dev_b = (not s["ld"]) and abs(cb.mean - s["cb"].mean) >= k * s["cb"].sigma_overall
            trows.append([d, n, _r(cw.mean), _r(cw.cp), _r(cw.cpk), "N/A" if s["ld"] else _r(cb.mean),
                          "N/A" if s["ld"] else _r(cb.cp), "N/A" if s["ld"] else _r(cb.cpk),
                          "⚠ 일자 평균 이탈" if (dev_w or dev_b) else "정상"])
            for j, v in ((3, cw.cp), (4, cw.cpk)) + (() if s["ld"] else ((6, cb.cp), (7, cb.cpk))):
                tf[(i, j)] = grade(v)[1]
            if dev_w or dev_b:
                tf[(i, 8)] = "FEF3C7"
        row = _table(ws3, row, 2, ["측정일", "n", "WU 평균", "WU Cp", "WU Cpk", "BU 평균", "BU Cp", "BU Cpk", "판정"], trows, tf) + 2

    # ---- 4. NG 이력
    ws4 = wb.create_sheet("NG 이력")
    ng_mask = (data["WU"] < SPECS["WU"][0]) | ((data["BU"] < SPECS["BU"][0]) & (data["LD"].astype(str) != "Y")) | \
              (data["판정"].astype(str).str.upper() == "NG")
    ng = data[ng_mask][["이벤트", "측정일", "LotID", "WU", "BU", "판정", "원본파일"]].sort_values(["이벤트", "측정일"])
    ws4.append(list(ng.columns))
    for rec in ng.itertuples(index=False):
        ws4.append([str(v) if isinstance(v, date) else v for v in rec])
    ws4.auto_filter.ref = f"A1:G{len(ng) + 1}"
    for col, w in zip("ABCDEFG", (14, 12, 48, 9, 9, 8, 46)):
        ws4.column_dimensions[col].width = w

    # ---- 5. Raw
    ws5 = wb.create_sheet("Raw")
    cols = ["이벤트", "매핑상태"] + [c for c in db.DATA_COLS if c in data.columns]
    ws5.append(cols)
    for rec in data[cols].itertuples(index=False):
        ws5.append([None if (isinstance(v, float) and np.isnan(v)) else (str(v) if isinstance(v, date) else v) for v in rec])
    ws5.auto_filter.ref = f"A1:{ws5.cell(1, len(cols)).column_letter}{len(data) + 1}"

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    log(f"✅ 리포트 저장: {Path(out_path).name}")
    return str(out_path)
