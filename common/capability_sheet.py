"""'공정능력_BU_WU' 시트 작성 (Tool A 결과 엑셀 안, D1 A안)."""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import charts
from capability import SPECS, capability, grade, is_invalid_wu

SHEET_NAME = "공정능력_BU_WU"
FONT = "맑은 고딕"
_thin = Side(style="thin", color="CBD5E1")
_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
_HEAD = PatternFill("solid", fgColor="1E3A5F")


def _table(ws, r0, c0, header, rows, fills=None):
    for j, h in enumerate(header):
        c = ws.cell(r0, c0 + j, h)
        c.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        c.fill = _HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BOX
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = ws.cell(r0 + i, c0 + j, v)
            c.font = Font(name=FONT, size=9, bold=(j == 0))
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = _BOX
            if fills and (i - 1, j) in fills:
                c.fill = PatternFill("solid", fgColor=fills[(i - 1, j)])
        ws.row_dimensions[r0 + i].height = 20


def _img(ws, png: bytes, anchor: str, width_px: int):
    img = XLImage(io.BytesIO(png))
    ratio = width_px / img.width
    img.width, img.height = width_px, int(img.height * ratio)
    ws.add_image(img, anchor)
    return int(img.height * 0.75 / 15) + 1          # 대략 차지하는 행 수(행 높이 15pt 기준)


def _r(v, nd=2):
    try:
        return round(float(v), nd) if pd.notna(v) else "—"
    except (TypeError, ValueError):
        return "—"


def compute(df: pd.DataFrame, ld: bool) -> dict:
    """df: 측정 순서로 정렬된 WU·BU 열. 반환 {'WU': Capability, 'BU': Capability}"""
    return {"WU": capability(df["WU"], "WU", raw_count=len(df)),
            "BU": capability(df["BU"], "BU", ld=ld, raw_count=len(df))}


def write_sheet(wb, df: pd.DataFrame, ld: bool, meta: dict) -> dict:
    """df 열: LotID, WU, BU, time(측정 시각, datetime). meta: model, source, version"""
    df = df.sort_values(["time", "LotID"], na_position="last").reset_index(drop=True)
    n_bad = int(is_invalid_wu(df["WU"]).sum())          # WU -100(측정 이상값) 행은 계산·그림에서 제외
    df = df[~is_invalid_wu(df["WU"])].reset_index(drop=True)
    caps = compute(df, ld)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMNOP", (2, 10, 12, 9, 9, 9, 7, 7, 8, 8, 8, 13, 13, 12, 13, 30)):
        ws.column_dimensions[col].width = w

    ws["B2"] = f"공정능력 분석 (BU/WU) — {meta.get('model', '')}".rstrip(" —")
    ws["B2"].font = Font(name=FONT, size=16, bold=True, color="1E3A5F")
    ws["B3"] = (f"{meta.get('source', '')} · LotID별 최신 1건 기준 n={len(df)} · σ(군내)=이동범위/1.128 · "
                f"{meta.get('version', '')} · {datetime.now():%Y-%m-%d %H:%M}")
    ws["B3"].font = Font(name=FONT, size=9, color="6B7280")
    if n_bad:
        ws["B4"] = f"⚠ WU -100(측정 이상값) {n_bad}건 제외 — Cp/Cpk·그림·층별표에 넣지 않음"
        ws["B4"].font = Font(name=FONT, size=10, bold=True, color="B91C1C")

    header = ["항목", "n (사용/제외)", "평균", "σ 군내", "σ 전체", "LSL", "USL", "Cp", "Cpk", "Ppk",
              "판정 (Cpk)", "실측 NG", "정규 가정 예측 NG", "정규성", "비고"]
    rows, fills = [], {}
    for i, name in enumerate(("WU", "BU")):
        c = caps[name]
        if c.ld:
            rows.append([name, f"0 / {c.n_excluded}", "—", "—", "—", c.lsl, c.usl, "N/A", "N/A", "N/A",
                         "N/A (LD 모델)", "—", "—", "—", "LD 모델 — BU 미측정(고정값)으로 제외"])
            for j in (7, 8, 9, 10, 14):
                fills[(i, j)] = "E5E7EB"
            continue
        rows.append([name, f"{c.n} / {c.n_excluded}", _r(c.mean), _r(c.sigma_within), _r(c.sigma_overall), c.lsl, c.usl,
                     _r(c.cp), _r(c.cpk), _r(c.ppk), c.verdict, f"{c.ng}ea ({c.ng_pct:.1f}%)",
                     f"{c.expected_ng_pct:.1f}%" if pd.notna(c.expected_ng_pct) else "—", c.normality,
                     " / ".join(c.warnings)])
        fills[(i, 7)] = grade(c.cp)[1]                 # D6 A안: Cp·Cpk 둘 다 색상
        fills[(i, 8)] = grade(c.cpk)[1]
        fills[(i, 10)] = grade(c.cpk)[1]
        if c.normality != "정규":
            fills[(i, 13)] = "FEF3C7"
    _table(ws, 5, 2, header, rows, fills)
    ws["B8"] = ("판정 기준: ≥1.67 양호 · 1.33~1.67 적합 · 1.00~1.33 주의 · <1.00 부적합 (Cp·Cpk 같은 기준) · "
                "USL 100은 균일도의 물리적 상한 · Ppk는 양식 19행 CPK와 같은 값")
    ws["B8"].font = Font(name=FONT, size=8, color="6B7280")

    row = 10
    dates = df["time"].dt.strftime("%m-%d").fillna("시각없음") if "time" in df else None
    wu = caps["WU"]
    h1 = _img(ws, charts.hist_with_spec(df["WU"].astype(float), "WU", *SPECS["WU"], wu), f"B{row}", 560) if wu.n >= 2 else 0
    bu = caps["BU"]
    if not bu.ld and bu.n >= 2:
        h2 = _img(ws, charts.hist_with_spec(df["BU"].astype(float), "BU", *SPECS["BU"], bu), f"I{row}", 560)
    else:
        ws[f"I{row + 2}"] = "BU: LD 모델로 계산 제외"
        ws[f"I{row + 2}"].font = Font(name=FONT, size=11, color="6B7280")
        h2 = 4
    row += max(h1, h2) + 1
    if wu.n >= 2:
        row += _img(ws, charts.ichart(df["WU"].astype(float), "WU", SPECS["WU"][0], dates), f"B{row}", 1130) + 1
    if not bu.ld and bu.n >= 2:
        row += _img(ws, charts.ichart(df["BU"].astype(float), "BU", SPECS["BU"][0], dates), f"B{row}", 1130) + 1

    ws[f"B{row}"] = "층별표 — 측정일별"
    ws[f"B{row}"].font = Font(name=FONT, size=10, bold=True, color="1E3A5F")
    body, f2 = [], {}
    for i, (d, g) in enumerate(df.groupby(df["time"].dt.date, sort=True)):
        cw = capability(g["WU"], "WU")
        cb = capability(g["BU"], "BU", ld=ld)
        body.append([str(d), len(g), _r(cw.mean), _r(cw.cp), _r(cw.cpk), f"{cw.ng}ea",
                     "—" if ld else _r(cb.mean), "N/A" if ld else _r(cb.cp), "N/A" if ld else _r(cb.cpk),
                     "—" if ld else f"{cb.ng}ea"])
        for j, v in ((3, cw.cp), (4, cw.cpk)) + (() if ld else ((7, cb.cp), (8, cb.cpk))):
            f2[(i, j)] = grade(v)[1]
    _table(ws, row + 1, 2, ["측정일", "n", "WU 평균", "WU Cp", "WU Cpk", "WU NG", "BU 평균", "BU Cp", "BU Cpk", "BU NG"], body, f2)

    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    return caps


def write_compare_sheet(wb, full: pd.DataFrame, sub_mask, ld: bool, meta: dict) -> dict:
    """(Tool B, 4번 ②안) 판정은 측정 CSV 전체 기준, 이미지 정리본은 참고로 함께.
    full 열: LotID, WU, BU, time(측정 시각). sub_mask: full과 같은 길이의 True/False (이미지 정리본 LotID)
    meta: model, csv_files, version. 반환 {'WU': (전체, 정리본), 'BU': (전체, 정리본)}"""
    full = full.assign(_sub=list(sub_mask)).sort_values(["time", "LotID"], na_position="last").reset_index(drop=True)
    n_bad = int(is_invalid_wu(full["WU"]).sum())        # WU -100(측정 이상값) 행은 계산·그림에서 제외
    full = full[~is_invalid_wu(full["WU"])].reset_index(drop=True)
    sub = full[full["_sub"]]
    caps = {"WU": (capability(full["WU"], "WU", raw_count=len(full)), capability(sub["WU"], "WU", raw_count=len(sub))),
            "BU": (capability(full["BU"], "BU", ld=ld, raw_count=len(full)), capability(sub["BU"], "BU", ld=ld, raw_count=len(sub)))}
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGHIJKLMNOPQ", (2, 8, 16, 11, 9, 9, 9, 7, 7, 8, 8, 8, 13, 13, 12, 13, 46)):
        ws.column_dimensions[col].width = w
    t = pd.to_datetime(full["time"], errors="coerce").dropna()
    period = "" if t.empty else (f"{t.min():%Y-%m-%d}" if t.min().date() == t.max().date()
                                 else f"{t.min():%Y-%m-%d} ~ {t.max():%Y-%m-%d}")
    ws["B2"] = f"공정능력 분석 (BU/WU) — {meta.get('model', '')}".rstrip(" —")
    ws["B2"].font = Font(name=FONT, size=16, bold=True, color="1E3A5F")
    ws["B3"] = (f"판정 = 이번 실행에서 고른 측정 CSV 전체(LotID별 최신 1건) · 이미지 정리본 = 결과 시트의 LotID(참고) · "
                f"측정 CSV {meta.get('csv_files', '?')}개 · 측정 기간 {period or '—'} · σ(군내)=이동범위/1.128(측정 시각 순) · "
                f"{meta.get('version', '')} · {datetime.now():%Y-%m-%d %H:%M}")
    ws["B3"].font = Font(name=FONT, size=9, color="6B7280")
    if n_bad:
        ws["B4"] = f"⚠ WU -100(측정 이상값) {n_bad}건 제외 — Cp/Cpk·그림에 넣지 않음"
        ws["B4"].font = Font(name=FONT, size=10, bold=True, color="B91C1C")

    header = ["항목", "대상", "n (사용/제외)", "평균", "σ 군내", "σ 전체", "LSL", "USL", "Cp", "Cpk", "Ppk",
              "판정 (Cpk)", "실측 NG", "정규 가정 예측 NG", "정규성", "비고"]
    rows, fills = [], {}
    for name in ("WU", "BU"):
        for k, (who, c) in enumerate((("측정 전체 (판정)", caps[name][0]), ("이미지 정리본 (참고)", caps[name][1]))):
            i = len(rows)
            if c.ld:
                rows.append([name, who, f"0 / {c.n_excluded}", "—", "—", "—", c.lsl, c.usl, "N/A", "N/A", "N/A",
                             "N/A (LD 모델)", "—", "—", "—", "LD 모델 — BU 미측정(고정값)으로 제외"])
                for j in (8, 9, 10, 11):
                    fills[(i, j)] = "E5E7EB"
                continue
            if c.n < 2:
                rows.append([name, who, f"{c.n} / {c.n_excluded}", "—", "—", "—", c.lsl, c.usl, "—", "—", "—", "—",
                             "—", "—", "—", "데이터 부족"])
                continue
            rows.append([name, who, f"{c.n} / {c.n_excluded}", _r(c.mean), _r(c.sigma_within), _r(c.sigma_overall), c.lsl, c.usl,
                         _r(c.cp), _r(c.cpk), _r(c.ppk), c.verdict if k == 0 else f"({c.verdict})",
                         f"{c.ng}ea ({c.ng_pct:.1f}%)",
                         f"{c.expected_ng_pct:.1f}%" if pd.notna(c.expected_ng_pct) else "—", c.normality,
                         " / ".join(c.warnings)])
            fills[(i, 8)] = grade(c.cp)[1]
            fills[(i, 9)] = grade(c.cpk)[1]
            if k == 0:
                fills[(i, 11)] = grade(c.cpk)[1]
            if c.normality != "정규":
                fills[(i, 14)] = "FEF3C7"
    _table(ws, 5, 2, header, rows, fills)
    ws["B10"] = ("판정 기준: ≥1.67 양호 · 1.33~1.67 적합 · 1.00~1.33 주의 · <1.00 부적합 (Cp·Cpk 같은 기준) · "
                 "USL 100은 균일도의 물리적 상한 · 정리본 값은 고른 LotID에 따라 달라지므로 참고용")
    ws["B10"].font = Font(name=FONT, size=8, color="6B7280")

    row = 12
    wa, ws_ = caps["WU"]
    h1 = _img(ws, charts.hist_overlay(full["WU"], sub["WU"], "WU", *SPECS["WU"], wa, ws_), f"B{row}", 600) if wa.n >= 2 else 0
    ba, bs = caps["BU"]
    if not ba.ld and ba.n >= 2:
        h2 = _img(ws, charts.hist_overlay(full["BU"], sub["BU"], "BU", *SPECS["BU"], ba, bs), f"J{row}", 600)
    else:
        ws[f"J{row + 2}"] = "BU: LD 모델로 계산 제외"
        ws[f"J{row + 2}"].font = Font(name=FONT, size=11, color="6B7280")
        h2 = 4
    row += max(h1, h2) + 1
    dates = pd.to_datetime(full["time"], errors="coerce").dt.strftime("%m-%d").fillna("시각없음")
    lab = f"이미지 정리본 {int(full['_sub'].sum())}ea"
    if wa.n >= 2:
        row += _img(ws, charts.ichart(full["WU"].astype(float), "WU", SPECS["WU"][0], dates, highlight=full["_sub"],
                                      highlight_label=lab), f"B{row}", 1180) + 1
    if not ba.ld and ba.n >= 2:
        row += _img(ws, charts.ichart(full["BU"].astype(float), "BU", SPECS["BU"][0], dates, highlight=full["_sub"],
                                      highlight_label=lab), f"B{row}", 1180) + 1
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    return caps
