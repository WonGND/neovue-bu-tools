"""공통 모듈 검증. 사용: python test_common.py <Daily data xlsx> <임시 폴더>
1) 공정능력: 첨부 WINE 775ea로 이전 계산값·양식 19행 수식(Ppk)과 일치
2) LD 감지 3) 누적 DB 중복 방지·이벤트 매핑"""
import sys
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "common"))
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

import history_db as db  # noqa: E402
from capability import capability, grade, is_constant  # noqa: E402

SRC, TMP = Path(sys.argv[1]), Path(sys.argv[2])
TMP.mkdir(parents=True, exist_ok=True)
ok = 0


def check(cond, msg):
    global ok
    assert cond, msg
    ok += 1
    print("  OK", msg)


ws = openpyxl.load_workbook(SRC, data_only=True).worksheets[0]
df = pd.DataFrame([[ws.cell(r, 3).value, ws.cell(r, 19).value, ws.cell(r, 20).value] for r in range(20, 795)],
                  columns=["LotID", "WU", "BU"])
df["k"] = df.LotID.str[32:38] + df.LotID.str[-5:]
df = df.sort_values("k")
print("[1] 공정능력 (WINE 775ea)")
wu, bu = capability(df.WU, "WU"), capability(df.BU, "BU")
for c, exp in ((wu, (2.05, 1.04, 0.88)), (bu, (1.07, 0.53, 0.45))):
    got = (round(c.cp, 2), round(c.cpk, 2), round(c.ppk, 2))
    check(got == exp, f"{c.name} Cp/Cpk/Ppk = {got} (기대 {exp})")
for name, lsl in (("WU", 80), ("BU", 50)):
    x = df[name].astype(float)
    template = (x.mean() - lsl) / (3 * x.std(ddof=1))        # 양식 19행: (AVERAGE-LSL)/(3*STDEV)
    c = capability(x, name)
    check(abs(template - c.ppk) < 1e-9, f"{name} Ppk = 양식 19행 수식 ({template:.4f})")
check(wu.ng == 10 and bu.ng == 44, f"실측 NG WU {wu.ng} / BU {bu.ng} (기대 10 / 44)")
check(grade(1.04)[0] == "주의" and grade(0.53)[0] == "부적합" and grade(2.05)[0] == "양호", "판정 등급")
check(any("비정규" in w for w in wu.warnings), "WU 비정규 경고")

print("[2] LD 감지")
check(is_constant([50.0] * 30), "BU 전부 50 → LD")
check(not is_constant(df.BU), "WINE BU → LD 아님")
ld = capability([50.0] * 30, "BU", ld=True)
check(ld.verdict == "N/A (LD 모델)" and ld.n == 0, "LD 모델 BU는 계산 제외")

print("[3] 누적 DB")
p = TMP / "test_db.xlsx"
p.unlink(missing_ok=True)
rows = pd.DataFrame({"LotID": ["GJ96111000-1032601120-VMA129BLAC01260415A00014", "X2"],
                     "측정일시": [datetime(2026, 4, 16, 8, 30), datetime(2026, 4, 16, 9, 0)],
                     "WU": [86.3, 85.1], "BU": [61.2, 55.0], "모델명": ["WINE", "WINE"]})
check(db.append(p, rows) == (2, 0), "첫 적재 2행")
check(db.append(p, rows) == (0, 2), "같은 파일 다시 적재 → 0행 (중복 건너뜀)")
data, events = db.load(p)
check(data.LotID.iloc[0] == "GJ961110001032601120VMA129BLAC01260415A00014", "LotID 정규화(하이픈 제거)")
check(db.prod_date(data.LotID.iloc[0]) == date(2026, 4, 15), "LotID 생산일 추출")
events = pd.DataFrame({"이벤트명": ["A", "B"], "모델": ["WINE", ""], "시작일": [date(2026, 4, 15), date(2026, 4, 16)],
                       "종료일": [date(2026, 4, 20), date(2026, 4, 16)], "메모": ["", ""]})
db.save_events(p, events)
data, events = db.load(p)
m = db.assign_events(data, events)
check(m["매핑상태"].iloc[0].startswith("겹침"), "두 이벤트 기간이 겹치면 '겹침'")
events.loc[1, "시작일"] = date(2026, 5, 1)
events.loc[1, "종료일"] = date(2026, 5, 2)
m = db.assign_events(data, events)
check(list(m["이벤트"]) == ["A", "A"], "기간 안 → 이벤트 A (v0.1 기간 이벤트)")
events.loc[0, "모델"] = "IND"
check(set(db.assign_events(data, events)["매핑상태"]) == {"미지정"}, "모델 불일치 → 미지정")

print("[4] v0.2 날짜 묶기 · 중복 방지 · 공용 DB")
events.loc[0, "모델"] = "WINE"
db.save_events(p, events)
data, events = db.load(p)
check(set(events["구분"]) == {"기간"}, "v0.1 이벤트는 '기간'으로 읽음")
data, events, n = db.migrate_period_events(data, events)
check(n == 2 and set(data[db.LABEL]) == {"A"} and set(events["구분"]) == {"묶음"}, "기간 이벤트 → 행에 이벤트 붙임(묶음 전환)")
db.save(p, data, events)
up = rows.assign(측정일시=[datetime(2026, 4, 16)] * 2, 출처=db.UPLOAD_SOURCE, **{db.LABEL: "B"})
st = db.append_rows(p, up, label_dups=True)
check(st["added"] == 0 and st["skipped"] == 2, "같은 측정(LotID·WU·BU)은 시각이 달라도 다시 넣지 않음")
check(st["relabeled"] == 0, "이미 이벤트가 있는 행은 다른 이벤트로 바꾸지 않음")
p2 = TMP / "test_db2.xlsx"
p2.unlink(missing_ok=True)
db.append_rows(p2, up.assign(**{db.LABEL: ""}))                    # 과거 파일 먼저 (미지정)
st = db.append_rows(p2, rows.assign(출처="Tool A v1.5"))            # 같은 측정이 Tool A로 들어옴
d2, _ = db.load(p2)
check(st["upgraded"] == 2 and len(d2) == 2 and d2["측정일시"].dt.hour.tolist() == [8, 9], "업로드분 → Tool A 측정 시각으로 보강")
st = db.append_rows(p2, up.assign(**{db.LABEL: "C"}), label_dups=True)
d2, e2 = db.load(p2)
check(st["relabeled"] == 2 and set(d2[db.LABEL]) == {"C"} and "C" in set(e2["이벤트명"]), "미지정 행은 파일 추가 때 고른 이벤트로 묶임")
key = db.day_key(d2).iloc[0]
check(db.set_label(p2, [("C", key)], "") == 2 and set(db.load(p2)[0][db.LABEL]) == {""}, "묶음 풀기 → 미지정")
check(db.set_label(p2, [("", key)], "D", "메모D") == 2, "날짜 묶기")
d2, e2 = db.load(p2)
r = e2[e2["이벤트명"] == "D"].iloc[0]
check(r["시작일"] == date(2026, 4, 16) and r["메모"] == "메모D" and r["모델"] == "WINE", "기간·모델은 데이터에서 자동")
db.rename_event(p2, "D", "D2", "새 메모")
check(set(db.load(p2)[0][db.LABEL]) == {"D2"}, "이름 바꾸면 행도 같이")
check(db.delete_rows(p2, [("D2", key)]) == 2 and len(db.load(p2)[0]) == 0, "날짜 행 삭제")
check(p2.with_name(p2.stem + ".backup.xlsx").exists(), "삭제 전 백업 파일")
# 공용 위치 합치기
import os
home = TMP / "home"
os.environ["HOME"] = str(home)
legacy = TMP / "exe_dir" / db.DB_NAME
legacy.parent.mkdir(parents=True, exist_ok=True)
legacy.unlink(missing_ok=True)
db.append_rows(legacy, rows.assign(출처="Tool A v1.5"))
shared = db.default_db_path()
shared.unlink(missing_ok=True)
merged = []
path, msg = db.adopt_shared_db(str(legacy), legacy, merged)
check(Path(path) == shared and len(db.load(shared)[0]) == 2 and legacy.exists(), "exe 옆 옛 DB → 문서\\BU_Tools로 합침(원본 유지)")
path, msg2 = db.adopt_shared_db(path, legacy, merged)
check(msg2 == "" and len(db.load(shared)[0]) == 2, "두 번째 실행은 다시 합치지 않음")
custom = str(TMP / "my.xlsx")
check(db.adopt_shared_db(custom, legacy, merged)[0] == custom, "직접 고른 DB 경로는 그대로")

print("[5] WU -100(측정 이상값) 제외 — Tool A·B 공정능력 시트")
import capability_sheet as cs  # noqa: E402
from capability import is_invalid_wu  # noqa: E402
rng = np.random.default_rng(1)
dfx = pd.DataFrame({"LotID": [f"L{i:03d}" for i in range(60)], "WU": rng.normal(88, 2, 60), "BU": rng.normal(70, 5, 60),
                    "time": pd.date_range("2026-10-01", periods=60, freq="min")})
dfx.loc[[5, 17, 40], "WU"] = -100
check(int(is_invalid_wu(dfx["WU"]).sum()) == 3, "-100 3건 감지")
wbx = openpyxl.Workbook()
capsx = cs.write_sheet(wbx, dfx, False, {"model": "TEST", "source": "t", "version": "Tool A test"})
clean = dfx[dfx["WU"] != -100]
check(capsx["WU"].n == 57 and abs(capsx["WU"].mean - clean["WU"].mean()) < 1e-9, "Tool A 시트: -100 행 빼고 계산 (n=57)")
check("WU -100" in str(wbx[cs.SHEET_NAME]["B4"].value), "Tool A 시트: 제외 알림")
wby = openpyxl.Workbook()
capsy = cs.write_compare_sheet(wby, dfx, [i % 2 == 0 for i in range(60)], False, {"model": "TEST", "csv_files": 1, "version": "Tool B test"})
check(capsy["WU"][0].n == 57 and "WU -100" in str(wby[cs.SHEET_NAME]["B4"].value), "Tool B 시트: -100 행 빼고 계산 + 알림")
print(f"전체 {ok}건 통과")
