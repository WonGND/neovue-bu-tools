"""트렌드 리포트 v0.6 로직 시험 (화면 없이)
사용: python run_tool_c.py <Tool A 시험 출력 폴더(DB 포함)> <리포트 출력 폴더>
1) 과거 파일(B, C) 추가 + 이벤트 지정 → 다시 추가해도 0건
2) Tool A가 이미 쌓은 날짜의 통합정리 파일을 추가하며 이벤트 지정 → 새로 넣지 않고 미지정 행만 묶음
3) 남은 미지정 날짜를 이벤트로 묶기 (G10은 LD 모델)
4) 리포트 생성 (기간은 데이터에서 자동) · 5) v0.3 변화 감지 시트(설명·기준·의미)"""
import glob
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "Trend"))
sys.path.insert(0, str(ROOT / "common"))
from openpyxl import load_workbook  # noqa: E402

import history_db as db  # noqa: E402
import trend_report as tr  # noqa: E402

A_OUT, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)
DB = A_OUT / "BU_History_DB.xlsx"


def add_files(files, label, dates=None, memo=None):
    """화면의 '파일 추가 → 이벤트 지정'과 같은 순서 (BU_trend_report.TrendApp._append)"""
    data, events = db.load(DB)
    total = {"added": 0, "skipped": 0, "relabeled": 0, "upgraded": 0}
    for i, f in enumerate(files):
        p = tr.parse_daily_file(f)
        d = (dates or {}).get(i, p["date"])
        data, st = db.merge_rows(data, tr.rows_for_db(p, d, f, label), label_dups=bool(label))
        for k in total:
            total[k] += st[k]
    db.save(DB, data, db.ensure_events(data, events, {label: memo} if (label and memo) else None))
    return total


def last(tag):
    return sorted(glob.glob(str(A_OUT / tag / "Daily data 통합정리_*.xlsx")))[-1]


d0, _ = db.load(DB)
n0 = len(d0)
print(f"시작: DB {n0}행 (Tool A 자동 적재분, 전부 미지정 {int((d0[db.LABEL] == '').sum())}행)")

# 1) 과거 파일 추가 + 이벤트 지정 (파일명 날짜는 생성일이라 측정일은 직접 지정)
for tag, d, ev in (("B_0610", date(2026, 6, 10), "이벤트 B"), ("C_0826", date(2026, 8, 26), "이벤트 C")):
    st = add_files([last(tag)], ev, {0: d}, memo="오후 측정분 재측정(예시)" if ev == "이벤트 C" else None)
    print(f"{tag} → {ev}: {st}")
    assert st["added"] > 0
    st = add_files([last(tag)], ev, {0: d})
    assert st["added"] == 0 and st["skipped"] > 0, "같은 파일 다시 추가는 0건"

# 2) Tool A가 이미 쌓은 날(04-21)의 통합정리 파일을 '이벤트 A'로 추가
st = add_files([last("A_0421")], "이벤트 A")
print("A_0421 통합정리 → 이벤트 A:", st)
assert st["added"] == 0 and st["relabeled"] == st["skipped"] > 0, "Tool A 적재분은 새로 넣지 않고 이벤트만 붙어야 함"

# 3) 남은 미지정 날짜 묶기
data, events = db.load(DB)
un = data[data[db.LABEL] == ""]
keys = sorted(set(db.day_key(un)))
print("미지정 날짜:", keys)
a_keys = [k for k in keys if k.startswith("2026-04")]          # 10-02는 G10(LD) 시험 데이터
g_keys = [k for k in keys if k not in a_keys]
db.set_label(DB, [("", k) for k in a_keys], "이벤트 A")
db.set_label(DB, [("", k) for k in g_keys], "G10 확인", "LD 모델")
data, events = db.load(DB)
assert (data[db.LABEL] != "").all(), "모든 행이 이벤트에 속해야 함"
print(events[["이벤트명", "모델", "시작일", "종료일", "메모", "구분"]].to_string(index=False))
print(data.groupby(db.LABEL).size().to_string())

# v0.5: 이벤트 C 3개 Lot ID의 WU를 -100(측정 이상값)으로 → 분석 제외 + 알람
data, events = db.load(DB)
c_idx = data.index[data[db.LABEL] == "이벤트 C"][:3]
bad_lots = set(data.loc[c_idx, "LotID"])
data.loc[c_idx, "WU"] = -100
db.save(DB, data, events)

# 4) 리포트
out = tr.generate(DB, ["이벤트 A", "이벤트 B", "이벤트 C", "G10 확인"], OUT / "트렌드리포트_test.xlsx", log=print)
wb = load_workbook(out)
print("시트:", wb.sheetnames)
ws = wb["이벤트 요약"]
h = next(r for r in range(1, 80) if ws.cell(r, 2).value == "이벤트" and ws.cell(r, 3).value == "기간")   # v0.4: 그림 아래로 내려감
for r in range(h + 1, h + 5):
    print([ws.cell(r, c).value for c in range(2, 13)])
assert ws.cell(h + 1, 3).value == "2026-04-16 ~ 2026-04-21", "기간은 데이터에서 자동 계산"
assert ws.cell(h + 4, 2).value == "G10 확인" and ws.cell(h + 4, 12).value == "N/A (LD)", "G10은 LD로 BU 제외"
# 5) 변화 감지 시트 (v0.5: '내용' 열, 경고 기준 바꾸기 표 없음, WU -100 알람)
assert wb.sheetnames[:2] == ["이벤트 요약", "변화 감지"], wb.sheetnames
wc = wb["변화 감지"]
hr = next(r for r in range(1, 30) if wc.cell(r, 2).value == "이벤트")
assert "WU -100" in str(wc.cell(10, 2).value) and "3건" in str(wc.cell(10, 2).value), wc.cell(10, 2).value
head = [wc.cell(hr, c).value for c in range(2, 10)]
assert head[-1] == "내용" and head[5] == "경고 기준", head
assert not any(wc.cell(r, 2).value == "경고 기준 바꾸기" for r in range(1, 80)), "경고 기준 바꾸기 표 제외"
rows = [[wc.cell(r, c).value for c in range(2, 10)] for r in range(hr + 1, hr + 40) if wc.cell(r, 2).value]
hits = [r for r in rows if str(r[6]).startswith("⚠")]
assert hits and all(r[7] and "→" not in r[7] for r in hits), "⚠ 행 내용 (화살표 안내 없음)"
assert all(r[1] for r in rows[:len(hits)]), "비교 대상(직전 이벤트) 표시"
link = [ws.cell(r, 4) for r in range(h, h + 15) if ws.cell(r, 4).hyperlink]
assert link and "확인 필요" in link[0].value and str(len(hits)) in link[0].value, "요약 시트에 건수 + 바로가기"
print(f"변화 감지 시트 OK — 확인 필요 {len(hits)}건 / {len(rows)}행 · 요약: {link[0].value}")
# 6) 배치: 요약 핵심 그림 2개 + 결론 메모 칸 · 미시 이벤트 접기 (v0.4)
assert len(ws._images) == 2, "요약 시트 핵심 그림 A·B"
assert any("결론 메모" == ws.cell(r, 2).value for r in range(1, 60)), "결론 메모 칸"
wm = wb["측정일별(미시)"]
levels = [r for r, d in wm.row_dimensions.items() if d.outlineLevel == 1]
assert levels and all(type(im.anchor).__name__ == "TwoCellAnchor" for im in wm._images), "행 그룹 + 그림이 행과 같이 접힘"
links = [c for c in wm[5] if c.hyperlink]
assert len(links) == 4, "미시 바로가기 4개"
# 7) v0.5: WU -100 제외 · 색좌표 확대 그림 · NG 이력 이미지 링크 없음
alarm = [ws.cell(r, 2).value for r in range(h, h + 8) if "WU -100" in str(ws.cell(r, 2).value)]
assert alarm and "이벤트 C 3건" in alarm[0], alarm
c_row = next(r for r in range(h + 1, h + 5) if ws.cell(r, 2).value == "이벤트 C")
assert "(−3)" in str(ws.cell(c_row, 5).value), ws.cell(c_row, 5).value
wn = wb["NG 이력"]
assert wn.max_column == 7 and not (set(str(wn.cell(r, 3).value) for r in range(2, wn.max_row + 1)) & bad_lots), "-100 행은 NG 아님"
wr = wb["Raw"]
col = [wr.cell(1, c).value for c in range(1, wr.max_column + 1)].index("분석 제외") + 1
assert sum(wr.cell(r, col).value == "WU -100" for r in range(2, wr.max_row + 1)) == 3, "Raw에 제외 표시"
d2, _ = db.load(DB)
g = d2[(d2[db.LABEL] == "이벤트 C") & (d2["WU"] != -100)].sort_values("측정일시").drop_duplicates("LotID", keep="last")
exp = round(float(g["WU"].mean()), 2)                  # 평균은 순서와 무관 → -100 빠졌는지 확인
assert abs(ws.cell(c_row, 6).value - exp) < 0.006 and len(g) == 809, (ws.cell(c_row, 6).value, exp, len(g))
assert len(wb["트렌드(거시)"]._images) == 8, "거시: 색좌표 전체 + 확대"
assert len(wm._images) == 4 * 3 + 4 * 2 - 1, "미시: 이벤트마다 색좌표 이동 2장 (G10은 BU 없음)"
print(f"WU -100 제외 OK — {alarm[0][:40]}… · 이벤트 C WU 평균 {exp} (n=809)")
# 8) v0.6: 진행률 콜백 · 휘도/CR 박스플롯 · 메모리 수정 함수
prog = []
tr.generate(DB, ["이벤트 A", "이벤트 B", "이벤트 C"], OUT / "트렌드리포트_WINE.xlsx", log=print,
            progress=lambda p_, t_: prog.append((p_, t_)))
pcts = [x[0] for x in prog]
assert pcts == sorted(pcts) and pcts[-1] == 100 and any("측정일별" in x[1] for x in prog), prog
print(f"진행률 {len(prog)}단계: " + " → ".join(f"{p_:.0f}%" for p_ in pcts))
d0, e0 = db.load(DB)
key = sorted(set(db.day_key(d0[d0[db.LABEL] == "이벤트 B"])))[0]
d1, e1, n = db.set_label_df(d0, e0, [("이벤트 B", key)], "이벤트 A")
assert n > 0 and (d1[db.LABEL] == "이벤트 B").sum() == 0 and (d0[db.LABEL] == "이벤트 B").sum() == n, "메모리 수정은 원본을 안 바꿈"
d2, e2 = db.rename_event_df(d1, e1, "이벤트 A", "이벤트 A2", "메모")
assert "이벤트 A2" in set(e2["이벤트명"]) and (d2[db.LABEL] == "이벤트 A2").sum() == (d1[db.LABEL] == "이벤트 A").sum()
print("트렌드 리포트 v0.6 시험 통과")
