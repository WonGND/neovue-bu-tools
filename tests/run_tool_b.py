"""Tool B v0.6 시험 (화면 없이): 시험용 BU/WU 이미지 + 측정 CSV → crop_report.xlsx
사용: python run_tool_b.py <testdata 폴더> <출력 폴더>
1) 결과 시트는 v0.5와 같은지 (v0.5 = git 이전 버전 코드로 같은 입력 실행해서 비교)
2) 공정능력 시트: 판정 = 측정 CSV 전체, 정리본 = 이미지 있는 LotID → Tool A와 같은 계산 값
3) LD 모델(G10): BU 제외"""
import csv
import random
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "BU_LotID"))
sys.path.insert(0, str(ROOT / "common"))
from openpyxl import load_workbook  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402

import BU_organize_one_click_v05 as tb  # noqa: E402
from capability import capability  # noqa: E402

TD, OUT = Path(sys.argv[1]), Path(sys.argv[2])
if OUT.exists():
    shutil.rmtree(OUT)
OUT.mkdir(parents=True)


def make_case(tag, n_img, seed=3):
    rows = list(csv.DictReader(open(TD / f"LMK6DataLog_{tag}.csv", encoding="utf-8-sig")))
    lots = sorted({r["Panel_ID"] for r in rows})
    random.seed(seed)
    pick = random.sample(lots, min(n_img, len(lots)))
    img_root, data_root = OUT / tag / "BU Image", OUT / tag / "data"
    data_root.mkdir(parents=True)
    shutil.copy(TD / f"LMK6DataLog_{tag}.csv", data_root / "LMK6DataLog.csv")
    for lot in pick:
        d = img_root / lot
        d.mkdir(parents=True)
        for kind in ("BU", "WU"):
            im = Image.new("RGB", (160, 100), "black")
            ImageDraw.Draw(im).rectangle([30, 20, 130, 80], fill=(90, 90, 90) if kind == "BU" else (230, 230, 230))
            im.save(d / f"{lot}_{kind}_1.png")
    return img_root, data_root, rows, pick


def result_sheet(xlsx):
    ws = load_workbook(xlsx)["결과"]
    return [tuple(c.value for c in r) for r in ws.iter_rows()], len(ws._images)


# 1) A_0421: 이미지 100개 / 측정 CSV 전체
img_root, data_root, rows, pick = make_case("A_0421", 100)
res = tb.run_pipeline(img_root, data_root, 12, 20)
x = Path(res["excel_path"])
wb = load_workbook(x)
print("시트:", wb.sheetnames)
assert wb.sheetnames == ["결과", "공정능력_BU_WU"]
ws = wb["공정능력_BU_WU"]
table = [[ws.cell(r, c).value for c in range(2, 13)] for r in range(6, 10)]
for t in table:
    print(t)
# 측정 CSV 전체(LotID별 최신)로 직접 계산한 값과 같은지
latest = {}
for r in rows:
    k = tb.normalize_lot_id(r["Panel_ID"])
    if k not in latest or r["Time"] > latest[k]["Time"]:
        latest[k] = r
srt = sorted(latest.values(), key=lambda r: (r["Time"], r["Panel_ID"]))
cw = capability([float(r["White_Uniformity"]) for r in srt], "WU")
cb = capability([float(r["Black_Uniformity"]) for r in srt], "BU")
assert table[0][1] == "측정 전체 (판정)" and table[0][2] == f"{len(latest)} / 0"
# 열: 0 항목 1 대상 2 n 3 평균 4 σ군내 5 σ전체 6 LSL 7 USL 8 Cp 9 Cpk 10 Ppk
assert abs(table[0][8] - round(cw.cp, 2)) < 1e-9 and abs(table[0][9] - round(cw.cpk, 2)) < 1e-9, "WU 전체 Cp/Cpk"
assert abs(table[2][8] - round(cb.cp, 2)) < 1e-9 and abs(table[2][9] - round(cb.cpk, 2)) < 1e-9, "BU 전체 Cp/Cpk"
assert table[1][2] == "100 / 0" and table[3][2] == "100 / 0", "정리본 n = 이미지 LotID 100"
sub = [r for r in srt if r["Panel_ID"] in set(pick)]
cbs = capability([float(r["Black_Uniformity"]) for r in sub], "BU")
assert abs(table[3][9] - round(cbs.cpk, 2)) < 1e-9, "BU 정리본 Cpk"
print(f"A_0421 OK — 전체 n={len(latest)} BU Cpk {cb.cpk:.2f} · 정리본 n=100 BU Cpk {cbs.cpk:.2f}")

# 결과 시트가 v0.5와 같은지 (git 저장소 안에서만: git의 v0.5 코드로 같은 입력 실행)
_rel = "BU_LotID/BU_organize_one_click_v05.py"
_intro = subprocess.run(["git", "-C", str(ROOT), "log", "--format=%H", "-S", "TOOL_VERSION", "--", _rel],
                        capture_output=True, text=True).stdout.split()        # v0.6 코드가 처음 들어간 커밋
src = subprocess.run(["git", "-C", str(ROOT), "show", f"{_intro[-1]}^:{_rel}" if _intro else f"HEAD:{_rel}"],
                     capture_output=True)
if src.returncode != 0 or b"TOOL_VERSION" in src.stdout:
    print("v0.5 비교 생략 (git 기록 없음)")
else:
    v05 = OUT / "v05_code"
    v05.mkdir()
    (v05 / "tb05.py").write_bytes(src.stdout)
    r05 = subprocess.run([sys.executable, "-c", f"""
import sys; sys.path.insert(0, r'{v05}')
from pathlib import Path
import tb05
from openpyxl import load_workbook
res = tb05.run_pipeline(Path(r'{img_root}'), Path(r'{data_root}'), 12, 20)
ws = load_workbook(res['excel_path'])['결과']
import json; print('JSON' + json.dumps([[str(c.value) for c in r] for r in ws.iter_rows()]))
print('IMGS', len(ws._images))
"""], capture_output=True, text=True, check=True).stdout
    import json  # noqa: E402
    old_rows = json.loads(r05.split("JSON", 1)[1].splitlines()[0])
    old_imgs = int(r05.split("IMGS", 1)[1].split()[0])
    res = tb.run_pipeline(img_root, data_root, 12, 20)     # v0.5 실행이 같은 경로에 다시 썼으므로 v0.6 결과를 다시 만듦
    new_rows, new_imgs = result_sheet(res["excel_path"])
    assert [[str(v) for v in r] for r in new_rows] == old_rows and new_imgs == old_imgs, "결과 시트가 v0.5와 다름"
    print(f"결과 시트 v0.5와 동일 ({len(old_rows)}행, 이미지 {old_imgs}개)")

# 2) 공정능력 시트 끄기
res = tb.run_pipeline(img_root, data_root, 12, 20, make_capability=False)
assert load_workbook(res["excel_path"]).sheetnames == ["결과"]
print("공정능력 시트 끄기 OK")

# 3) LD 모델
img_root, data_root, rows, pick = make_case("G10_LD", 20)
res = tb.run_pipeline(img_root, data_root, 12, 20)
ws = load_workbook(res["excel_path"])["공정능력_BU_WU"]
assert ws.cell(8, 13).value == "N/A (LD 모델)", ws.cell(8, 13).value      # M열 = 판정
print("G10 LD OK — BU N/A")
print("Tool B v0.6 시험 통과")
