import json
import os
import sys
import threading
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext

import pandas as pd
from openpyxl import load_workbook

# 공용 모듈(common/) 경로: 소스 실행 시 ../common, exe는 spec에서 함께 묶음
_HERE = Path(__file__).resolve().parent
for _p in (_HERE.parent / "common", _HERE / "common"):
    if _p.exists() and str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import capability_sheet  # noqa: E402
import history_db  # noqa: E402
from capability import is_constant, summary_line  # noqa: E402

VERSION = "v1.5"
APP_TITLE = f"One Click 데이터 정리 도구 {VERSION}"

# [설정] 시간 포맷 및 기본 양식 경로
TIME_FMT = "%Y.%m.%d %H:%M:%S"
DEFAULT_TEMPLATE_PATH = r"../../업무/00. 곽원용/01. PJT/2025/3. 12.9 inch_VDA129BLAC01/0. 양산 대응 폴더/Daily data 정리/양식_Daily data_BU설비_광학데이터_날짜.xlsx"
SUPPORTED_TIME_FORMATS = (
    "%Y.%m.%d %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
)
SETTINGS_NAME = "OneClick_settings.json"
DB_DEFAULT_NAME = "BU_History_DB.xlsx"

# 시트별 선별 컬럼 정의
CA_SELECTED_COLS = [
    "Time", "Panel", "Panel_ID",
    "Red_X", "Red_Y",
    "Green_X", "Green_Y",
    "Blue_X", "Blue_Y",
    "White_SX", "White_SY", "White_LV",
    "Black_LV", "EEPROM_59"
]

LMK_SELECTED_COLS = [
    "Time", "Panel", "Panel_ID", "Judge",
    "Black_Uniformity", "White_Uniformity"
]

TEMPLATE_START_ROW = 20
TEMPLATE_MAPPING = {
    'Red_X': 6, 'Red_Y': 7, 'Green_X': 8, 'Green_Y': 9, 'Blue_X': 10, 'Blue_Y': 11,
    'White_SX': 13, 'White_SY': 14, 'White_LV': 15, 'Black_LV': 16,
    'White_Uniformity': 19, 'Black_Uniformity': 20, 'EEPROM_59': 27
}


def app_base_dir() -> Path:
    return Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else _HERE


# ---------------------------------------------------------------- 데이터 처리 (화면과 분리)
def read_df(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == '.csv':
        last_error = None
        for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
            try:
                return pd.read_csv(path, encoding=encoding)
            except UnicodeDecodeError as exc:
                last_error = exc
        raise last_error
    return pd.read_excel(path)


def clean_df(df):
    df = df.copy()
    # 시간 처리
    if "Time" in df.columns:
        df["Time_dt"] = pd.to_datetime(df["Time"], format=TIME_FMT, errors="coerce")
        for fmt in SUPPORTED_TIME_FORMATS[1:]:
            missing = df["Time_dt"].isna()
            if not missing.any():
                break
            parsed = pd.to_datetime(df.loc[missing, "Time"], format=fmt, errors="coerce")
            df.loc[missing, "Time_dt"] = parsed
        df = df[df["Time_dt"].notna()]
        df = df.sort_values("Time_dt")

    # Panel_ID 기준 최신화
    if "Panel_ID" in df.columns:
        df["Panel_ID"] = df["Panel_ID"].astype(str).str.strip().str.replace(" ", "", regex=False).str.replace("-", "", regex=False).str.replace("_", "", regex=False)
        df["Panel_ID"] = df["Panel_ID"].str.replace(".0", "", regex=False).str.upper()
        df = df.drop_duplicates(subset=["Panel_ID"], keep="last")
        df = df.sort_values("Panel_ID")

    return df


def detect_ld(lmk_path) -> bool:
    """LMK6 DataLog의 BU(Black_Uniformity)가 전부 같은 값이면 LD 모델로 판단 (D3 B안)."""
    if not lmk_path or not os.path.exists(lmk_path):
        return False
    try:
        df = read_df(lmk_path)
    except Exception:
        return False
    return "Black_Uniformity" in df.columns and is_constant(df["Black_Uniformity"])


def save_individual_file(df, mode, save_dir):
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    out_path = os.path.join(save_dir, f"{mode}_cleaned_{timestamp}.xlsx")
    selected_cols = CA_SELECTED_COLS if mode == "CA410" else LMK_SELECTED_COLS
    actual_cols = [c for c in selected_cols if c in df.columns]   # 실제 데이터에 있는 컬럼만 (에러 방지)
    df_selected = df[actual_cols].copy()
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Sheet1_All", index=False)
        df_selected.to_excel(writer, sheet_name="Sheet2_Selected", index=False)
    return out_path


def _template_meta(ws):
    cut = ws["R2"].value if ws["R2"].value else ""
    code = ws["R3"].value if ws["R3"].value else ""
    return str(cut).strip(), str(code).strip()


def _cap_frame(df):
    """공정능력 계산용 표: LotID, WU, BU, time(LMK 측정 시각 우선)"""
    t = df["LMK_Time"] if "LMK_Time" in df else df.get("Time_dt")
    if "Time_dt" in df and t is not None:
        t = t.fillna(df["Time_dt"])
    return pd.DataFrame({"LotID": df["Panel_ID"].values,
                         "WU": pd.to_numeric(df.get("White_Uniformity"), errors="coerce").values,
                         "BU": pd.to_numeric(df.get("Black_Uniformity"), errors="coerce").values,
                         "time": pd.to_datetime(t).values})


def _db_rows(df, cut_model, ld, ca_path, lmk_path):
    t = df["LMK_Time"] if "LMK_Time" in df else df.get("Time_dt")
    if "Time_dt" in df and t is not None:
        t = t.fillna(df["Time_dt"])
    num = lambda c: pd.to_numeric(df[c], errors="coerce") if c in df else None
    rows = pd.DataFrame({
        "LotID": df["Panel_ID"].values,
        "모델명": cut_model,
        "측정일시": pd.to_datetime(t).values,
        "채널": df["Panel"].values if "Panel" in df else None,
        "Red_X": num("Red_X"), "Red_Y": num("Red_Y"), "Green_X": num("Green_X"), "Green_Y": num("Green_Y"),
        "Blue_X": num("Blue_X"), "Blue_Y": num("Blue_Y"),
        "White_x": num("White_SX"), "White_y": num("White_SY"),
        "White_LV": num("White_LV"), "Black_LV": num("Black_LV"),
        "WU": num("White_Uniformity"), "BU": num("Black_Uniformity"),
        "판정": df["Judge"].values if "Judge" in df else None,
        "EEPROM_59": df["EEPROM_59"].astype(str).values if "EEPROM_59" in df else None,
    })
    rows["CR"] = pd.to_numeric(rows["White_LV"], errors="coerce") / pd.to_numeric(rows["Black_LV"], errors="coerce")
    rows["모델코드"] = rows["LotID"].map(history_db.model_code)
    rows["생산일"] = rows["LotID"].map(history_db.prod_date)
    rows["LD"] = "Y" if ld else ""
    rows["출처"] = f"Tool A {VERSION}"
    rows["원본파일"] = " / ".join(os.path.basename(p) for p in (ca_path, lmk_path) if p)
    rows["Tool"] = f"OneClick {VERSION}"
    return rows


def process(opts: dict, log=print) -> dict:
    """opts: ca_path, lmk_path, template_path, save_dir, make_individual, make_cpk, ld, append_db, db_path"""
    ca_path, lmk_path = opts.get("ca_path") or "", opts.get("lmk_path") or ""
    save_dir, temp_path = opts["save_dir"], opts.get("template_path") or ""
    ld = bool(opts.get("ld"))
    result = {"out_path": None, "caps": None, "db": None}
    df_ca = df_lmk = None
    if ca_path:
        log("CA410 데이터 정리 중...")
        raw = read_df(ca_path)
        df_ca = clean_df(raw)
        log(f"   CA410 {len(raw)}건 → {len(df_ca)}건 (중복 {len(raw) - len(df_ca)}건 제거)")
    if lmk_path:
        log("LMK6 데이터 정리 중...")
        raw = read_df(lmk_path)
        df_lmk = clean_df(raw)
        log(f"   LMK6 {len(raw)}건 → {len(df_lmk)}건 (중복 {len(raw) - len(df_lmk)}건 제거)")

    lmk_file = None
    if opts.get("make_individual"):
        if df_ca is not None:
            log("CA410 개별 정리 파일(2개 시트) 생성 중...")
            res = save_individual_file(df_ca, "CA410", save_dir)
            if not (ca_path and lmk_path):
                result["out_path"] = res
        if df_lmk is not None:
            log("LMK6 개별 정리 파일(2개 시트) 생성 중...")
            lmk_file = save_individual_file(df_lmk, "LMK6", save_dir)
            if not (ca_path and lmk_path):
                result["out_path"] = lmk_file

    merged, wb, cut_model, code = None, None, "", ""
    if ca_path and lmk_path:
        if not os.path.exists(temp_path):
            raise FileNotFoundError(f"양식 파일을 찾을 수 없습니다: {temp_path}")
        if "Panel_ID" not in df_ca.columns or "Panel_ID" not in df_lmk.columns:
            raise ValueError("CA410/LMK6 양쪽 모두 Panel_ID 컬럼이 필요합니다.")
        log("데이터 병합 중...")
        lmk_subset = df_lmk[[c for c in ["Panel_ID", "Black_Uniformity", "White_Uniformity", "Judge", "Time_dt"]
                             if c in df_lmk.columns]].rename(columns={"Time_dt": "LMK_Time"})
        merged = pd.merge(df_ca, lmk_subset, on="Panel_ID", how="outer").sort_values("Panel_ID")
        log("양식 기입 중...")
        wb = load_workbook(temp_path)
        ws = wb.active
        cut_model, code = _template_meta(ws)
        for i, row in enumerate(merged.itertuples(), start=TEMPLATE_START_ROW):
            ws.cell(row=i, column=2, value=i - TEMPLATE_START_ROW + 1)
            ws.cell(row=i, column=3, value=row.Panel_ID)
            for attr, col in TEMPLATE_MAPPING.items():
                if hasattr(row, attr):
                    ws.cell(row=i, column=col, value=getattr(row, attr))

    # ---- v1.5: 공정능력 시트 (BU/WU, LMK6 데이터가 있을 때)
    base_df = merged if merged is not None else df_lmk
    if opts.get("make_cpk") and base_df is not None and "White_Uniformity" in base_df.columns:
        log(f"공정능력 계산 중... (LotID별 최신 1건, n = {len(base_df)}){' · LD 모델: BU 제외' if ld else ''}")
        model = " ".join(x for x in (cut_model, code) if x) or history_db.model_code(base_df["Panel_ID"].iloc[0])
        meta = {"model": model, "version": f"Tool A {VERSION}",
                "source": " / ".join(os.path.basename(p) for p in (ca_path, lmk_path) if p)}
        target_wb = wb
        if target_wb is None:
            if lmk_file:
                target_wb = load_workbook(lmk_file)
            else:
                from openpyxl import Workbook
                target_wb = Workbook()
                target_wb.remove(target_wb.active)
        caps = capability_sheet.write_sheet(target_wb, _cap_frame(base_df), ld, meta)
        result["caps"] = caps
        for c in caps.values():
            log("   " + summary_line(c))
            for w in c.warnings:
                if not c.ld:
                    log(f"   ⚠ {c.name}: {w}")
        if wb is None:
            if lmk_file:
                target_wb.save(lmk_file)
            else:
                result["out_path"] = os.path.join(save_dir, f"공정능력_BU_WU_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
                target_wb.save(result["out_path"])

    if wb is not None:
        out_path = os.path.join(save_dir, f"Daily data 통합정리_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        wb.save(out_path)
        result["out_path"] = out_path
        log(f"✅ 통합 완료: {os.path.basename(out_path)}")
        if opts.get("make_cpk") and result["caps"]:
            log(f"   └ 시트 추가: {capability_sheet.SHEET_NAME}")

    # ---- v1.5: 누적 DB 적재 (트렌드 리포트용, 측정 시각 포함)
    if opts.get("append_db") and base_df is not None:
        try:
            rows = _db_rows(base_df, cut_model, ld, ca_path, lmk_path)
            added, skipped = history_db.append(opts["db_path"], rows)
            result["db"] = (added, skipped)
            log(f"누적 DB 적재: {added}행 추가" + (f", 이미 있는 {skipped}행 건너뜀" if skipped else ""))
        except history_db.DBLockedError as exc:
            log(f"⚠ 누적 DB 적재 실패: {exc}")
    return result


# ---------------------------------------------------------------- 화면
class OneClickApp:
    def __init__(self, root):
        self.root = root
        self.root.title(APP_TITLE)
        self.root.geometry("820x900")

        self.last_saved_file = None
        self.last_saved_folder = None
        self.template_path = self.get_default_template_path()
        self._main_thread = threading.current_thread()
        self.settings_path = app_base_dir() / SETTINGS_NAME
        self.settings = self._load_settings()

        self.setup_ui()
        self.log("프로그램이 시작되었습니다.")
        self.check_template()

    # ---- 설정 저장 (각자 PC)
    def _load_settings(self):
        s = {"db_path": str(app_base_dir() / DB_DEFAULT_NAME), "make_cpk": True, "append_db": True}
        try:
            s.update(json.loads(self.settings_path.read_text(encoding="utf-8")))
        except (OSError, ValueError):
            pass
        return s

    def _save_settings(self):
        self.settings.update({"db_path": self.ent_db.get().strip(), "make_cpk": self.var_cpk.get(),
                              "append_db": self.var_db.get()})
        try:
            self.settings_path.write_text(json.dumps(self.settings, ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError:
            pass

    def setup_ui(self):
        # 파일 선택 프레임
        file_frame = tk.LabelFrame(self.root, text="입력 파일 설정 (CSV/Excel)", padx=10, pady=10)
        file_frame.pack(fill="x", padx=10, pady=5)

        tk.Label(file_frame, text="1. CA410 (색좌표):").grid(row=0, column=0, sticky="w")
        self.ent_ca = tk.Entry(file_frame)
        self.ent_ca.grid(row=0, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="찾기", command=lambda: self.browse_file(self.ent_ca, "CA410")).grid(row=0, column=2)

        tk.Label(file_frame, text="2. LMK6 (BU/WU):").grid(row=1, column=0, sticky="w", pady=5)
        self.ent_lmk = tk.Entry(file_frame)
        self.ent_lmk.grid(row=1, column=1, sticky="ew", padx=5)
        self.ent_lmk.bind("<FocusOut>", lambda _e: self.update_ld_detection())
        tk.Button(file_frame, text="찾기", command=lambda: self.browse_file(self.ent_lmk, "LMK6")).grid(row=1, column=2)

        tk.Label(file_frame, text="3. 양식 파일:").grid(row=2, column=0, sticky="w")
        self.ent_temp = tk.Entry(file_frame)
        self.ent_temp.insert(0, self.template_path if self.template_path else "")
        self.ent_temp.grid(row=2, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="변경", command=self.browse_template).grid(row=2, column=2)

        tk.Label(file_frame, text="4. 저장 폴더:").grid(row=3, column=0, sticky="w", pady=5)
        self.ent_save = tk.Entry(file_frame)
        self.ent_save.grid(row=3, column=1, sticky="ew", padx=5)
        tk.Button(file_frame, text="찾기", command=self.browse_folder).grid(row=3, column=2)

        self.var_make_individual = tk.BooleanVar(value=True)
        self.chk_individual = tk.Checkbutton(file_frame, text="개별 정리 파일(Cleaned) 생성 (2개 시트 구성)", variable=self.var_make_individual)
        self.chk_individual.grid(row=4, column=0, columnspan=3, sticky="w", pady=5)
        file_frame.columnconfigure(1, weight=1)

        # v1.5 분석 옵션
        opt = tk.LabelFrame(self.root, text="분석 옵션", padx=10, pady=8)
        opt.pack(fill="x", padx=10, pady=5)
        self.var_cpk = tk.BooleanVar(value=bool(self.settings.get("make_cpk", True)))
        tk.Checkbutton(opt, text="공정능력(Cp/Cpk) 시트 생성  —  WU 80~100 · BU 50~100", variable=self.var_cpk).grid(row=0, column=0, columnspan=3, sticky="w")
        self.var_ld = tk.BooleanVar(value=False)
        tk.Checkbutton(opt, text="LD 모델 (BU 미측정)  —  BU Cp/Cpk 제외, 'N/A (LD 모델)' 표기", variable=self.var_ld,
                       command=self._hide_ld_note).grid(row=1, column=0, columnspan=3, sticky="w")
        self.lbl_ld = tk.Label(opt, text="", fg="#b45309", bg="#fef3c7", anchor="w", padx=6)
        self.var_db = tk.BooleanVar(value=bool(self.settings.get("append_db", True)))
        tk.Checkbutton(opt, text="누적 DB에 적재 (트렌드 리포트용, 측정 시각 함께 저장)", variable=self.var_db).grid(row=3, column=0, columnspan=3, sticky="w")
        tk.Label(opt, text="   DB 위치:").grid(row=4, column=0, sticky="w")
        self.ent_db = tk.Entry(opt)
        self.ent_db.insert(0, self.settings.get("db_path", ""))
        self.ent_db.grid(row=4, column=1, sticky="ew", padx=5)
        tk.Button(opt, text="변경", command=self.browse_db).grid(row=4, column=2)
        opt.columnconfigure(1, weight=1)

        # 실행 버튼
        self.btn_run = tk.Button(self.root, text="▶ One-Click 정리 시작",
                                 bg="#28a745", fg="white", font=("맑은 고딕", 12, "bold"),
                                 height=2, command=self.start_process)
        self.btn_run.pack(fill="x", padx=10, pady=10)

        # 로그창
        log_frame = tk.LabelFrame(self.root, text="처리 로그", padx=10, pady=10)
        log_frame.pack(fill="both", expand=True, padx=10, pady=5)
        self.txt_log = scrolledtext.ScrolledText(log_frame, height=15, state="disabled", font=("Consolas", 10))
        self.txt_log.pack(fill="both", expand=True)

        # 결과 버튼
        res_frame = tk.Frame(self.root)
        res_frame.pack(fill="x", padx=10, pady=10)
        self.btn_open_f = tk.Button(res_frame, text="결과 파일 열기", state="disabled", command=self.open_last_file)
        self.btn_open_f.pack(side="left", expand=True, fill="x", padx=5)
        self.btn_open_d = tk.Button(res_frame, text="폴더 열기", state="disabled", command=self.open_last_folder)
        self.btn_open_d.pack(side="left", expand=True, fill="x", padx=5)

    # ---- LD 자동 감지 (D3 B안: 데이터 검사로 미리 체크, 사람은 확인)
    def update_ld_detection(self):
        path = self.ent_lmk.get().strip()
        if detect_ld(path):
            self.var_ld.set(True)
            self.lbl_ld.config(text="ⓘ BU 값이 전부 같은 값입니다 (σ = 0) → LD 모델로 자동 선택했습니다. 아니면 체크를 해제하세요.")
            self.lbl_ld.grid(row=2, column=0, columnspan=3, sticky="ew", pady=(2, 4))
            self.log("데이터 검사: BU 값이 모두 같음 → LD 모델로 자동 선택")
        else:
            self._hide_ld_note()

    def _hide_ld_note(self):
        self.lbl_ld.grid_remove()

    def log(self, msg):
        line = f"[{datetime.now().strftime('%H:%M:%S')}] {msg}\n"
        if threading.current_thread() is self._main_thread:
            self._append_log(line)
        else:
            self.root.after(0, self._append_log, line)

    def _append_log(self, line):
        self.txt_log.config(state="normal")
        self.txt_log.insert(tk.END, line)
        self.txt_log.see(tk.END)
        self.txt_log.config(state="disabled")

    def _set_run_button(self, state, text=None):
        self.btn_run.config(state=state, text=text if text is not None else self.btn_run.cget("text"))

    def _set_result_buttons(self, enabled: bool):
        state = "normal" if enabled else "disabled"
        self.btn_open_f.config(state=state)
        self.btn_open_d.config(state=state)

    def get_default_template_path(self):
        base = str(app_base_dir())
        candidate = os.path.normpath(os.path.join(base, DEFAULT_TEMPLATE_PATH))
        if os.path.exists(candidate):
            return candidate
        alt = os.path.normpath(os.path.join(os.path.dirname(base), DEFAULT_TEMPLATE_PATH))
        if os.path.exists(alt):
            return alt
        return candidate

    def check_template(self):
        path = self.ent_temp.get()
        if os.path.exists(path):
            self.log(f"양식 파일 감지됨: {os.path.basename(path)}")
            if not self.ent_save.get():
                self.ent_save.insert(0, os.path.dirname(path))
        else:
            self.log("⚠ 기본 경로에 양식 파일이 없습니다. 직접 선택해 주세요.")

    def browse_file(self, entry, name):
        p = filedialog.askopenfilename(title=f"{name} 파일 선택", filetypes=[("Data files", "*.csv *.xlsx *.xls")])
        if p:
            entry.delete(0, tk.END)
            entry.insert(0, p)
            if not self.ent_save.get():
                self.ent_save.insert(0, os.path.dirname(p))
            if entry is self.ent_lmk:
                self.update_ld_detection()

    def browse_template(self):
        p = filedialog.askopenfilename(title="양식 파일 선택", filetypes=[("Excel files", "*.xlsx")])
        if p:
            self.ent_temp.delete(0, tk.END)
            self.ent_temp.insert(0, p)

    def browse_folder(self):
        p = filedialog.askdirectory(title="저장 폴더 선택")
        if p:
            self.ent_save.delete(0, tk.END)
            self.ent_save.insert(0, p)

    def browse_db(self):
        p = filedialog.asksaveasfilename(title="누적 DB 파일 위치", defaultextension=".xlsx",
                                         initialfile=DB_DEFAULT_NAME, filetypes=[("Excel files", "*.xlsx")])
        if p:
            self.ent_db.delete(0, tk.END)
            self.ent_db.insert(0, p)

    def start_process(self):
        # v1.5: 입력 검사를 화면 스레드에서 먼저 (v1.4.1의 '처리 중...' 버튼 고착 수정)
        ca_path, lmk_path = self.ent_ca.get().strip(), self.ent_lmk.get().strip()
        if not any([ca_path, lmk_path]):
            messagebox.showwarning("알림", "정리할 파일을 하나 이상 선택해 주세요.")
            return
        if not self.ent_save.get().strip():
            messagebox.showwarning("알림", "저장 폴더를 지정해 주세요.")
            return
        if self.var_cpk.get() and lmk_path and not self.var_ld.get() and detect_ld(lmk_path):
            if messagebox.askyesno("LD 모델 확인", "BU 값이 모두 같습니다 (σ = 0).\nLD 모델(BU 미측정)로 보입니다.\n\nBU Cp/Cpk를 계산에서 제외할까요?"):
                self.var_ld.set(True)
        self._save_settings()
        opts = {"ca_path": ca_path, "lmk_path": lmk_path, "template_path": self.ent_temp.get().strip(),
                "save_dir": self.ent_save.get().strip(), "make_individual": self.var_make_individual.get(),
                "make_cpk": self.var_cpk.get(), "ld": self.var_ld.get(),
                "append_db": self.var_db.get(), "db_path": self.ent_db.get().strip()}
        self.btn_run.config(state="disabled", text="처리 중...")
        threading.Thread(target=self.run, args=(opts,), daemon=True).start()

    def run(self, opts):
        try:
            result = process(opts, self.log)
            out_path = result.get("out_path")
            if out_path:
                self.last_saved_file = out_path
                self.last_saved_folder = opts["save_dir"]
                self.root.after(0, self._set_result_buttons, True)
                self.root.after(0, messagebox.showinfo, "성공", "데이터 정리가 완료되었습니다.")
            else:
                self.log("⚠ 작업은 완료되었으나 생성된 파일이 없습니다 (옵션 확인).")
        except Exception as e:
            self.log(f"❌ 에러: {str(e)}")
            self.root.after(0, messagebox.showerror, "오류", f"처리 중 에러 발생:\n{str(e)}")
        finally:
            self.root.after(0, self._set_run_button, "normal", "▶ One-Click 정리 시작")

    def open_last_file(self):
        if self.last_saved_file:
            os.startfile(self.last_saved_file)

    def open_last_folder(self):
        if self.last_saved_folder:
            os.startfile(self.last_saved_folder)


if __name__ == "__main__":
    root = tk.Tk()
    app = OneClickApp(root)
    root.mainloop()
